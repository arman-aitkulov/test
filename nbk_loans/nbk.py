"""Загрузка и разбор статистики Нацбанка РК «Кредиты банковского сектора в региональном разрезе»."""
from __future__ import annotations

import io
import re
from datetime import date

import openpyxl
import pandas as pd
import requests

BASE = "https://nationalbank.kz"
RUBRIC_URL = BASE + "/ru/news/banking-sector-loans-to-economy-analytics/rubrics/{rubric}"
# Страница текущего года; с неё берутся ссылки на страницы остальных лет.
START_RUBRIC = 2602
FILE_TITLE = "кредиты банковского сектора в региональном разрезе"

SHEETS = {
    "Выдано": "Выдано кредитов за месяц",
    "Остатки": "Остаток задолженности на начало месяца",
    "Просрочка": "Просроченная задолженность на начало месяца",
}

MONTHS = {
    "январ": 1, "феврал": 2, "март": 3, "апрел": 4, "ма": 5, "июн": 6,
    "июл": 7, "август": 8, "сентябр": 9, "октябр": 10, "ноябр": 11, "декабр": 12,
}

_session = requests.Session()
_session.headers["User-Agent"] = "Mozilla/5.0 (nbk-loans-dashboard)"


def _get(url: str) -> requests.Response:
    resp = _session.get(url, timeout=60)
    resp.raise_for_status()
    return resp


def list_year_rubrics() -> dict[int, int]:
    """{год: id рубрики} по ссылкам-переключателям лет на странице раздела."""
    html = _get(RUBRIC_URL.format(rubric=START_RUBRIC)).text
    found = re.findall(
        r'banking-sector-loans-to-economy-analytics/rubrics/(\d+)"[^>]*>\s*(\d{4})\s*<', html
    )
    return {int(year): int(rubric) for rubric, year in found}


def find_regional_file(rubric: int) -> str | None:
    """Ссылка на xlsx «Кредиты банковского сектора в региональном разрезе» на странице года."""
    html = _get(RUBRIC_URL.format(rubric=rubric)).text
    for href, title in re.findall(
        r'href="(/file/download/\d+)"[^>]*download>(.*?)</a>', html, re.S
    ):
        text = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", title)).strip().lower()
        if text.startswith(FILE_TITLE):
            return BASE + href
    return None


def _parse_period(header: str) -> date | None:
    """«за январь 2026 года» / «на 1 февраля 2026 года1» -> date(2026, 1|2, 1)."""
    header = header.lower()
    year = re.search(r"(20\d{2})", header)
    if not year:
        return None
    for stem, month in MONTHS.items():
        # «ма» — май/мая; проверяем после «март», поэтому ищем целым словом
        pattern = r"\bма[йя]\b" if stem == "ма" else r"\b" + stem
        if re.search(pattern, header):
            return date(int(year.group(1)), month, 1)
    return None


def _parse_sheet(ws) -> pd.DataFrame:
    rows = list(ws.iter_rows(values_only=True))
    # Строка с заголовками периодов — первая, где есть «20XX года».
    hdr_idx = next(
        i for i, r in enumerate(rows)
        if any(isinstance(v, str) and re.search(r"20\d{2}", v) for v in r[1:])
    )
    periods = {
        col: p for col, v in enumerate(rows[hdr_idx])
        if isinstance(v, str) and (p := _parse_period(v))
    }
    records = []
    # Берём только первый блок (все кредиты) — до строки «в том числе:».
    for r in rows[hdr_idx + 1:]:
        label = r[0]
        if label is None:
            if any(isinstance(v, str) and "в том числе" in v for v in r):
                break
            continue
        region = re.sub(r"\s+", " ", str(label)).strip()
        if region.lower().startswith("всего"):
            continue
        for col, period in periods.items():
            val = r[col] if col < len(r) else None
            if isinstance(val, (int, float)):
                records.append({"region": region, "period": period, "value": float(val)})
    return pd.DataFrame(records)


def parse_workbook(content: bytes) -> pd.DataFrame:
    """Возвращает long-таблицу: sheet, region, period, value (млн тенге)."""
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    frames = []
    for name in wb.sheetnames:
        key = next((k for k in SHEETS if name.strip().lower().startswith(k.lower())), None)
        if key is None:
            continue
        df = _parse_sheet(wb[name])
        df["sheet"] = key
        frames.append(df)
    return pd.concat(frames, ignore_index=True)


def load_year(rubric: int) -> pd.DataFrame:
    url = find_regional_file(rubric)
    if url is None:
        return pd.DataFrame(columns=["sheet", "region", "period", "value"])
    return parse_workbook(_get(url).content)


def load_all() -> pd.DataFrame:
    frames = [load_year(rubric) for rubric in list_year_rubrics().values()]
    df = pd.concat(frames, ignore_index=True)
    return df.drop_duplicates(["sheet", "region", "period"]).sort_values(["sheet", "period", "region"])
